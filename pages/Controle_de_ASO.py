import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import os

# --- AJUSTE DE CAMINHO PARA GITHUB ---
# Isso identifica a pasta onde o script está rodando
BASE_PATH = os.path.dirname(__file__) if "__file__" in locals() else "."

# Se o arquivo estiver dentro de uma pasta 'pages', volta um nível para achar o logo
if "pages" in BASE_PATH:
    ROOT_PATH = os.path.abspath(os.path.join(BASE_PATH, ".."))
else:
    ROOT_PATH = BASE_PATH

# O logo deve estar dentro de uma pasta chamada 'assets' no GitHub
PATH_LOGO = os.path.join(ROOT_PATH, "assets", "adivitta.png")

# Configurações de Planilha (Mantidas)
SHEET_ID = "1G_oVT9gK-n_jGh5R4g65qUwK_MfQGvCX-SA4NHNNflU"
UNIDADES = {
    "CURITIBA": "145843404",
    "SÃO JOSÉ": "1537243911",
    "ITAJAÍ": "517454238",
    "JOINVILLE": "1940989392"
}

def load_data(gid):
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid={gid}"
    return pd.read_csv(url)

def gerar_docx(dados, tipo, data_sugestao):
    doc = Document()
    
    # 1. Inserir Logotipo (Caminho relativo ao GitHub)
    if os.path.exists(PATH_LOGO):
        section = doc.sections[0]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(PATH_LOGO, width=Inches(1.5))
    
    # 2. Título do Formulário
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_titulo = titulo.add_run(f"FORMULÁRIO PARA AGENDAMENTO {tipo}")
    run_titulo.bold = True
    run_titulo.font.size = Pt(14)

    # 3. Tabela de Dados
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    
    labels = ["Nome Completo", "Cargo", "Setor", "Unidade", "Cliente", "Local", "Data Sugestão"]
    valores = [
        str(dados['Nome']), 
        str(dados['Cargo']), 
        str(dados['Setor']), 
        st.session_state.get('unidade_nome', 'N/A'),
        "MACROMAQ", 
        "Arapoti",
        data_sugestao.strftime('%d/%m/%Y')
    ]
    
    for i, (label, val) in enumerate(zip(labels, valores)):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = val
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

    target = BytesIO()
    doc.save(target)
    return target.getvalue()

# --- Interface Streamlit ---
st.title("🛡️ Sistema de Controle de ASO")

try:
    aba_nome = st.sidebar.selectbox("Selecione a Unidade", list(UNIDADES.keys()))
    st.session_state.unidade_nome = aba_nome
    df = load_data(UNIDADES[aba_nome])

    if not df.empty:
        hoje = datetime.now()
        prazo_limite = hoje + timedelta(days=10)
        df['Venc'] = pd.to_datetime(df['Venc'], dayfirst=True, errors='coerce')
        alertas = df[df['Venc'] <= prazo_limite].copy()
        
        st.metric("ASOs em Alerta (10 dias)", len(alertas))
        st.subheader(f"⚠️ Colaboradores em Alerta - {aba_nome}")
        st.dataframe(alertas[['Nome', 'Cargo', 'Setor', 'Venc']], use_container_width=True)

        st.markdown("---")
        st.subheader("📝 Gerar Solicitação de Agendamento")
        
        c1, c2, c3 = st.columns(3)
        with c1:
            nome_sel = st.selectbox("Colaborador", alertas['Nome'].tolist())
        with c2:
            tipo_ag = st.selectbox("Tipo", ["PERIÓDICO", "MUDANÇA DE RISCO", "RETORNO AO TRABALHO"])
        with c3:
            data_sug = st.date_input("Data de Sugestão", value=hoje + timedelta(days=2))

        if nome_sel:
            dados_colab = alertas[alertas['Nome'] == nome_sel].iloc[0]
            docx_file = gerar_docx(dados_colab, tipo_ag, data_sug)
            
            st.download_button(
                label=f"📥 Baixar Solicitação - {nome_sel}",
                data=docx_file,
                file_name=f"Solicitacao_{nome_sel.replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

except Exception as e:
    st.error(f"Erro no sistema: {e}")
